import os
import re

def rewrite():
    with open("build_report.py", "r", encoding="utf-8") as f:
        content = f.read()

    new_content = """import os
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def apply_margins(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)
    section.top_margin = Cm(2.5)
    section.right_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)

def initialize_document():
    document = Document()
    
    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(0)
    
    for section in document.sections:
        apply_margins(section)
        
    return document

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_cover_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SUMMER TRAINING REPORT")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 139)
    
    doc.add_paragraph()
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("ADVANCED ANDROID REVERSE ENGINEERING & AI-ASSISTED FORENSIC TRIAGE AUTOMATION")
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(16)
    run2.font.bold = True
    
    for _ in range(2): doc.add_paragraph()
    
    p3 = doc.add_paragraph("Submitted in partial fulfilment of the requirements\\nfor the award of the degree")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = p4.add_run("Bachelor of Technology\\nin\\nArtificial Intelligence & Machine Learning")
    run4.font.bold = True
    
    for _ in range(2): doc.add_paragraph()
    
    p5 = doc.add_paragraph()
    run5 = p5.add_run("Submitted by")
    run5.font.size = Pt(14)
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p6 = doc.add_paragraph("Name: Sahil Yadav\\nUniversity Roll No.: ____________\\nSemester: ____________\\nAcademic Year: 2026-27")
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    for _ in range(2): doc.add_paragraph()
    
    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run7 = p7.add_run("VIVEKANANDA INSTITUTE OF PROFESSIONAL STUDIES - TECHNICAL CAMPUS\\nGrade A++ Accredited Institution by NAAC")
    run7.font.bold = True
    
    p8 = doc.add_paragraph("NBA Accredited for MCA Programme; Recognized under Section 2(f) by UGC;\\nAffiliated to GGSIP University, Delhi; Recognized by Bar Council of India and AICTE\\nAn ISO 9001:2015 Certified Institution")
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER

def start_front_matter(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_margins(section)
    section.footer.is_linked_to_previous = False
    
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), 'lowerRoman')
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)
    
    # We do NOT add the page number to the footer here. It will just be accounted for.

def add_declaration(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DECLARATION")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    doc.add_paragraph()
    
    p2 = doc.add_paragraph('I hereby declare that the summer training report entitled "ADVANCED ANDROID REVERSE ENGINEERING & AI-ASSISTED FORENSIC TRIAGE AUTOMATION" is an authentic record of my own work for the requirements of summer training during the period from ________ to ________ for the award of degree of B.Tech. (AI&ML) from School of Engineering and Technology, Vivekananda Institute of Professional Studies – Technical Campus, Pitampura, New Delhi.')
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    for _ in range(4): doc.add_paragraph()
    doc.add_paragraph("(Signature of student)\\n(Name of Student: Sahil Yadav)\\n(University Roll No.: _________)\\nDate: ____________________")
    doc.add_page_break()

def add_image_page(doc, title, image_path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    doc.add_paragraph()
    
    if os.path.exists(image_path):
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run().add_picture(image_path, width=Cm(15))
    else:
        doc.add_paragraph(f"[Image {image_path} not found]")
        
    doc.add_page_break()

def add_acknowledgement(doc):
    # Start a new section for Acknowledgement to start printing the page numbers
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_margins(section)
    section.footer.is_linked_to_previous = False
    
    # Add page number to this footer so it prints from Acknowledgement onwards
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_para.add_run())
    
    for _ in range(2): doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ACKNOWLEDGEMENT")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    for _ in range(4): doc.add_paragraph()
    
    ack_text = [
        "I extend my deepest gratitude and sincere appreciation to the senior supervisory officers and technical mentors at the Intelligence Fusion and Strategic Operations (IFSO) Unit, Special Cell, Delhi Police, for their continuous guidance, invaluable mentorship, and support throughout the tenure of my internship.",
        "I am profoundly grateful for being granted the privilege to work within advanced forensic environments, leverage specialized analytical infrastructure, and contribute to real-world, high-impact investigative casework. Under the dedicated direction of my mentors, I gained deep technical exposure to mobile application triage, static and dynamic malware disassembly, threat intelligence correlation, and digital forensic methodologies.",
        "Crucially, it was through their expert insight and hands-on direction that I was introduced to advanced open-source reverse-engineering frameworks and developer tools, including JADX and extensible AI-driven workflows. Their guidance not only broadened my technical perspective on modern threat analysis and automation but also taught me how to effectively apply these open-source tools to solve complex cybercrime challenges. The knowledge, problem-solving mindset, and investigative rigor imparted by my mentors have been central to the successful conceptualization, engineering, and completion of the technical deliverables presented in this report."
    ]
    for para in ack_text:
        p2 = doc.add_paragraph(para)
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()
    p3 = doc.add_paragraph("Sahil Yadav\\nIntern")
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_page_break()

def add_seq_field(run, seq_identifier):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = f' SEQ {seq_identifier} \\* ARABIC '
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_field(paragraph, instruction_text):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = instruction_text
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_toc_lof_lot(doc):
    p_toc_title = doc.add_paragraph()
    p_toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_title.paragraph_format.space_before = Pt(80) # Approx 25-35mm more to reach 50-60mm from top margin
    p_toc_title.paragraph_format.space_after = Pt(72)  # 5-6 single line spacing from the title
    
    run = p_toc_title.add_run("TABLE OF CONTENTS")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    
    p_toc = doc.add_paragraph()
    add_field(p_toc, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()
    
    p_lof_title = doc.add_paragraph()
    p_lof_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lof_title.paragraph_format.space_before = Pt(80)
    run = p_lof_title.add_run("LIST OF FIGURES")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    p_lof = doc.add_paragraph()
    add_field(p_lof, 'TOC \\h \\z \\c "Figure"')
    doc.add_page_break()
    
    p_lot_title = doc.add_paragraph()
    p_lot_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lot_title.paragraph_format.space_before = Pt(80)
    run = p_lot_title.add_run("LIST OF TABLES")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    p_lot = doc.add_paragraph()
    add_field(p_lot, 'TOC \\h \\z \\c "Table"')
    doc.add_page_break()

def start_chapter(doc, is_first_chapter=False):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_margins(section)
    section.footer.is_linked_to_previous = False
    section.different_first_page_header_footer = True
    
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), 'decimal')
    if is_first_chapter:
        pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)
    
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_para.add_run())

"""
    
    # We replace from the top up to `def _parse_inline_formatting`
    content = re.sub(r'^.*?def _parse_inline_formatting', new_content + '\\ndef _parse_inline_formatting', content, flags=re.DOTALL)
    
    with open("build_report.py", "w", encoding="utf-8") as f:
        f.write(content)

if __name__ == "__main__":
    rewrite()
