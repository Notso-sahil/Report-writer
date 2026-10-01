from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, ns
from docx.oxml.ns import qn

def initialize_document():
    document = Document()
    
    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(0)
    
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3.5)
        section.top_margin = Cm(2.5)
        section.right_margin = Cm(1.25)
        section.bottom_margin = Cm(1.25)
        
    return document

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_cover_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SUMMER TRAINING REPORT")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 139)
    
    doc.add_paragraph()
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("ADVANCED ANDROID REVERSE ENGINEERING & AI-ASSISTED FORENSIC TRIAGE AUTOMATION")
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(16)
    run2.font.bold = True
    
    for _ in range(2): doc.add_paragraph()
    
    p3 = doc.add_paragraph("Submitted in partial fulfilment of the requirements\nfor the award of the degree")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = p4.add_run("Bachelor of Technology\nin\nArtificial Intelligence & Machine Learning")
    run4.font.bold = True
    
    for _ in range(2): doc.add_paragraph()
    
    p5 = doc.add_paragraph()
    run5 = p5.add_run("Submitted by")
    run5.font.size = Pt(14)
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p6 = doc.add_paragraph("Name: Sahil Yadav\nUniversity Roll No.: ____________\nSemester: ____________\nAcademic Year: 2026-27")
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    for _ in range(2): doc.add_paragraph()
    
    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run7 = p7.add_run("VIVEKANANDA INSTITUTE OF PROFESSIONAL STUDIES - TECHNICAL CAMPUS\nGrade A++ Accredited Institution by NAAC")
    run7.font.bold = True
    p8 = doc.add_paragraph("NBA Accredited for MCA Programme; Recognized under Section 2(f) by UGC;\nAffiliated to GGSIP University, Delhi; Recognized by Bar Council of India and AICTE\nAn ISO 9001:2015 Certified Institution")
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # No page break here, it's handled by start_front_matter

def start_front_matter(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    # Unlink footer from previous (Cover Page)
    section.footer.is_linked_to_previous = False
    
    # Set Roman pagination starting at i (1)
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), 'lowerRoman')
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)
    
    # Add page number to footer
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_para.add_run())

def add_declaration(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DECLARATION")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    doc.add_paragraph()
    
    p2 = doc.add_paragraph('I hereby declare that the summer training report entitled "ADVANCED ANDROID REVERSE ENGINEERING & AI-ASSISTED FORENSIC TRIAGE AUTOMATION" is an authentic record of my own work for the requirements of summer training during the period from ________ to ________ for the award of degree of B.Tech. (AI&ML) from School of Engineering and Technology, Vivekananda Institute of Professional Studies – Technical Campus, Pitampura, New Delhi.')
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    for _ in range(4): doc.add_paragraph()
    doc.add_paragraph("(Signature of student)\n(Name of Student: Sahil Yadav)\n(University Roll No.: _________)\nDate: ____________________")
    doc.add_page_break()

def add_image_page(doc, title, image_path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    doc.add_paragraph()
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run().add_picture(image_path, width=Cm(15))
    doc.add_page_break()

def add_acknowledgement(doc):
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ACKNOWLEDGEMENT")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    for _ in range(5): doc.add_paragraph()
    
    ack_text = [
        "I extend my deepest gratitude and sincere appreciation to the senior supervisory officers and technical mentors at the Intelligence Fusion and Strategic Operations (IFSO) Unit, Special Cell, Delhi Police, for their continuous guidance, invaluable mentorship, and support throughout the tenure of my internship.",
        "I am profoundly grateful for being granted the privilege to work within advanced forensic environments, leverage specialized analytical infrastructure, and contribute to real-world, high-impact investigative casework. Under the dedicated direction of my mentors, I gained deep technical exposure to mobile application triage, static and dynamic malware disassembly, threat intelligence correlation, and digital forensic methodologies.",
        "Crucially, it was through their expert insight and hands-on direction that I was introduced to advanced open-source reverse-engineering frameworks and developer tools, including JADX and extensible AI-driven workflows. Their guidance not only broadened my technical perspective on modern threat analysis and automation but also taught me how to effectively apply these open-source tools to solve complex cybercrime challenges. The knowledge, problem-solving mindset, and investigative rigor imparted by my mentors have been central to the successful conceptualization, engineering, and completion of the technical deliverables presented in this report."
    ]
    for para in ack_text:
        p2 = doc.add_paragraph(para)
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()
    p3 = doc.add_paragraph("Sahil Yadav\nIntern")
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_page_break()

def add_field(paragraph, instruction_text):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = instruction_text
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_toc_lof_lot(doc):
    # TOC
    p_toc_title = doc.add_paragraph()
    p_toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_toc_title.add_run("TABLE OF CONTENTS")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    
    p_toc = doc.add_paragraph()
    add_field(p_toc, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()
    
    # LOF
    p_lof_title = doc.add_paragraph()
    p_lof_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_lof_title.add_run("LIST OF FIGURES")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    
    p_lof = doc.add_paragraph()
    add_field(p_lof, 'TOC \\h \\z \\c "Figure"')
    doc.add_page_break()
    
    # LOT
    p_lot_title = doc.add_paragraph()
    p_lot_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_lot_title.add_run("LIST OF TABLES")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    
    p_lot = doc.add_paragraph()
    add_field(p_lot, 'TOC \\h \\z \\c "Table"')
    doc.add_page_break()
def start_chapter(doc, chapter_title, is_first_chapter=False):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.footer.is_linked_to_previous = False
    
    # Hide page number on the first page of this chapter
    section.different_first_page_header_footer = True
    
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), 'decimal')
    if is_first_chapter:
        pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)
    
    # Add page number to the normal footer (which applies to page 2 onwards)
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_para.add_run())
    
    # The first_page_footer is left empty intentionally.
    
    # Add Chapter Title
    p = doc.add_paragraph(style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(chapter_title)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

import re

def parse_markdown_to_docx(doc, md_text):
    lines = md_text.split('\n')
    in_code_block = False
    in_table = False
    table_data = []
    code_content = []
    
    for i, line in enumerate(lines):
        line = line.strip()
        
        # Code Blocks
        if line.startswith('```'):
            if in_code_block:
                p = doc.add_paragraph('\n'.join(code_content))
                p.style = doc.styles['No Spacing']
                p.paragraph_format.left_indent = Cm(1)
                run = p.runs[0]
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
                in_code_block = False
                code_content = []
                doc.add_paragraph()
            else:
                in_code_block = True
                code_content = []
            continue
            
        if in_code_block:
            code_content.append(line)
            continue
            
        if not line:
            continue
            
        # Table Captions
        if line.startswith('[TABLE_CAPTION]'):
            caption_text = line.replace('[TABLE_CAPTION]', '').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(caption_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            p.paragraph_format.space_after = Pt(6)
            continue
            
        # Tables
        if line.startswith('|') and line.endswith('|'):
            # Skip the separator line (e.g. |---|---|)
            if re.match(r'^\|[\s\-\|]+\|$', line):
                continue
                
            cols = [col.strip() for col in line.strip('|').split('|')]
            if not in_table:
                in_table = True
                table_data = [cols]
            else:
                table_data.append(cols)
                
            # If next line is not a table row, render the table
            if i + 1 == len(lines) or not lines[i+1].strip().startswith('|'):
                in_table = False
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for r_idx, row in enumerate(table_data):
                    for c_idx, cell_text in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_text
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.name = 'Times New Roman'
                                run.font.size = Pt(11)
                doc.add_paragraph() # Space after table
            continue
            
        # Figures
        if line.startswith('[FIGURE]'):
            img_path = line.replace('[FIGURE]', '').strip()
            import os
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img_path, width=Cm(14))
            continue
            
        # Figure Captions (BELOW)
        if line.startswith('[FIGURE_CAPTION]'):
            caption_text = line.replace('[FIGURE_CAPTION]', '').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(caption_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
            continue
            
        # Headings
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('#').strip()
            p = doc.add_paragraph(style=f'Heading {level}')
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            if level == 1:
                run.font.size = Pt(16)
            elif level == 2:
                run.font.size = Pt(14)
            else:
                run.font.size = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            continue
            
        # Bullet points
        if line.startswith('- ') or line.startswith('* '):
            text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            _parse_inline_formatting(p, text)
            continue
            
        # Blockquotes
        if line.startswith('> '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.runs = _parse_inline_formatting(p, text)
            continue
            
        # Regular paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _parse_inline_formatting(p, line)

def _parse_inline_formatting(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.font.bold = True
        else:
            paragraph.add_run(part)
    for run in paragraph.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)


if __name__ == "__main__":
    doc = initialize_document()
    
    # Section 1: Cover
    add_cover_page(doc)
    
    # Section 2: Front Matter (Roman)
    start_front_matter(doc)
    add_declaration(doc)
    add_image_page(doc, "OFFER LETTER", "offer_letter.png")
    add_image_page(doc, "CERTIFICATE", "Ifso_certificate.png")
    add_image_page(doc, "EMPLOYER FEEDBACK", "Employer Feedback Form.png")
    add_acknowledgement(doc)
    add_toc_lof_lot(doc)
    
    # Section 3: Chapter 1 (Arabic)
    start_chapter(doc, "Chapter-1. Introduction", is_first_chapter=True)
    with open("chapter_1.md", "r", encoding="utf-8") as f:
        parse_markdown_to_docx(doc, f.read())
    
    # Section 4: Chapter 2 (Arabic continuous)
    start_chapter(doc, "Chapter-2. Project Title", is_first_chapter=False)
    with open("chapter_2.md", "r", encoding="utf-8") as f:
        parse_markdown_to_docx(doc, f.read())
        
    # Section 5: Chapter 3 (Arabic continuous)
    start_chapter(doc, "Chapter-3. Case Study 1: Forensic Triage", is_first_chapter=False)
    with open("chapter_3.md", "r", encoding="utf-8") as f:
        parse_markdown_to_docx(doc, f.read())
        
    # Section 6: Chapter 4 (Arabic continuous)
    start_chapter(doc, "Chapter-4. Case Study 2: JADX-AI-MCP", is_first_chapter=False)
    with open("chapter_4.md", "r", encoding="utf-8") as f:
        parse_markdown_to_docx(doc, f.read())
        
    # Section 7: Chapter 5 (Arabic continuous)
    start_chapter(doc, "Chapter-5. Discussion & Future Scope", is_first_chapter=False)
    with open("chapter_5.md", "r", encoding="utf-8") as f:
        parse_markdown_to_docx(doc, f.read())
        
    doc.save("IFSO_Final_Report.docx")
    print("[+] Final Compilation generated successfully.")
