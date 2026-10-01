import os
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def apply_margins(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)
    section.top_margin = Cm(2.5)
    section.right_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)

def initialize_document():
    document = Document()
    
    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(0)
    
    for section in document.sections:
        apply_margins(section)
        
    return document

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_cover_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SUMMER TRAINING REPORT")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 139)
    
    doc.add_paragraph()
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("ADVANCED ANDROID REVERSE ENGINEERING & AI-ASSISTED FORENSIC TRIAGE AUTOMATION")
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(16)
    run2.font.bold = True
    
    for _ in range(2): doc.add_paragraph()
    
    p3 = doc.add_paragraph("Submitted in partial fulfilment of the requirements\nfor the award of the degree")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = p4.add_run("Bachelor of Technology\nin\nArtificial Intelligence & Machine Learning")
    run4.font.bold = True
    
    for _ in range(2): doc.add_paragraph()
    
    p5 = doc.add_paragraph()
    run5 = p5.add_run("Submitted by")
    run5.font.size = Pt(14)
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p6 = doc.add_paragraph("Name: Sahil Yadav\nUniversity Roll No.: ____________\nSemester: ____________\nAcademic Year: 2026-27")
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    for _ in range(2): doc.add_paragraph()
    
    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run7 = p7.add_run("VIVEKANANDA INSTITUTE OF PROFESSIONAL STUDIES - TECHNICAL CAMPUS\nGrade A++ Accredited Institution by NAAC")
    run7.font.bold = True
    
    p8 = doc.add_paragraph("NBA Accredited for MCA Programme; Recognized under Section 2(f) by UGC;\nAffiliated to GGSIP University, Delhi; Recognized by Bar Council of India and AICTE\nAn ISO 9001:2015 Certified Institution")
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER

def start_front_matter(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_margins(section)
    section.footer.is_linked_to_previous = False
    
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), 'lowerRoman')
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)

def add_declaration(doc):
    p = doc.add_paragraph(style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DECLARATION")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph()
    
    p2 = doc.add_paragraph('I hereby declare that the summer training report entitled "ADVANCED ANDROID REVERSE ENGINEERING & AI-ASSISTED FORENSIC TRIAGE AUTOMATION" is an authentic record of my own work for the requirements of summer training during the period from ________ to ________ for the award of degree of B.Tech. (AI&ML) from School of Engineering and Technology, Vivekananda Institute of Professional Studies – Technical Campus, Pitampura, New Delhi.')
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    for _ in range(4): doc.add_paragraph()
    doc.add_paragraph("(Signature of student)\n(Name of Student: Sahil Yadav)\n(University Roll No.: _________)\nDate: ____________________")
    doc.add_page_break()

def add_image_page(doc, title, image_path):
    p = doc.add_paragraph(style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph()
    
    if os.path.exists(image_path):
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run().add_picture(image_path, width=Cm(15))
    else:
        doc.add_paragraph(f"[Image {image_path} not found]")
        
    doc.add_page_break()

def add_acknowledgement(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_margins(section)
    section.footer.is_linked_to_previous = False
    
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_para.add_run())
    
    p = doc.add_paragraph(style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(72)
    run = p.add_run("ACKNOWLEDGEMENT")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    ack_text = [
        "I extend my deepest gratitude and sincere appreciation to the senior supervisory officers and technical mentors at the Intelligence Fusion and Strategic Operations (IFSO) Unit, Special Cell, Delhi Police, for their continuous guidance, invaluable mentorship, and support throughout the tenure of my internship.",
        "I am profoundly grateful for being granted the privilege to work within advanced forensic environments, leverage specialized analytical infrastructure, and contribute to real-world, high-impact investigative casework. Under the dedicated direction of my mentors, I gained deep technical exposure to mobile application triage, static and dynamic malware disassembly, threat intelligence correlation, and digital forensic methodologies.",
        "Crucially, it was through their expert insight and hands-on direction that I was introduced to advanced open-source reverse-engineering frameworks and developer tools, including JADX and extensible AI-driven workflows. Their guidance not only broadened my technical perspective on modern threat analysis and automation but also taught me how to effectively apply these open-source tools to solve complex cybercrime challenges. The knowledge, problem-solving mindset, and investigative rigor imparted by my mentors have been central to the successful conceptualization, engineering, and completion of the technical deliverables presented in this report."
    ]
    for para in ack_text:
        p2 = doc.add_paragraph(para)
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()
    p3 = doc.add_paragraph("Sahil Yadav\nIntern")
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_page_break()

def add_seq_field(run, seq_identifier):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = f' SEQ {seq_identifier} \\* ARABIC '
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_field(paragraph, instruction_text):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = instruction_text
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_toc_lof_lot(doc):
    p_toc_title = doc.add_paragraph()
    p_toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_title.paragraph_format.space_before = Pt(80)
    p_toc_title.paragraph_format.space_after = Pt(72)
    run = p_toc_title.add_run("TABLE OF CONTENTS")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    p_toc = doc.add_paragraph()
    add_field(p_toc, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()
    
    p_lof_title = doc.add_paragraph()
    p_lof_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lof_title.paragraph_format.space_before = Pt(80)
    run = p_lof_title.add_run("LIST OF FIGURES")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    p_lof = doc.add_paragraph()
    add_field(p_lof, 'TOC \\h \\z \\c "Figure"')
    doc.add_page_break()
    
    p_lot_title = doc.add_paragraph()
    p_lot_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lot_title.paragraph_format.space_before = Pt(80)
    run = p_lot_title.add_run("LIST OF TABLES")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    p_lot = doc.add_paragraph()
    add_field(p_lot, 'TOC \\h \\z \\c "Table"')
    doc.add_page_break()

def start_chapter(doc, is_first_chapter=False):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_margins(section)
    section.footer.is_linked_to_previous = False
    section.different_first_page_header_footer = True
    
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), 'decimal')
    if is_first_chapter:
        pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)
    
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_para.add_run())

def _parse_inline_formatting(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.font.bold = True
            run.font.name = 'Times New Roman'
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Times New Roman'
    for run in paragraph.runs:
        if run.font.name != 'Courier New':
            run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

def get_blocks(lines):
    blocks = []
    current_block = []
    block_type = "TEXT"
    in_code = False
    
    for line in lines:
        stripped = line.strip()
        
        if stripped.startswith('```'):
            if current_block:
                blocks.append((block_type, current_block))
                current_block = []
            in_code = not in_code
            block_type = "CODE" if in_code else "TEXT"
            continue
            
        if in_code:
            current_block.append(line)
            continue
            
        if not stripped:
            if current_block:
                blocks.append((block_type, current_block))
                current_block = []
            continue
            
        is_new_block = (
            stripped.startswith('#') or 
            stripped.startswith('![') or 
            stripped.startswith('|') or
            stripped == '<!-- -->' or
            re.match(r'^[\-\*]\s', stripped) or
            re.match(r'^\d+\.\s', stripped) or
            stripped.startswith('> ')
        )
        
        if is_new_block:
            if current_block:
                blocks.append((block_type, current_block))
                current_block = []
            block_type = "TABLE" if stripped.startswith('|') else "TEXT"
            current_block.append(stripped)
        else:
            if block_type == "TABLE" and not stripped.startswith('|'):
                blocks.append((block_type, current_block))
                current_block = [stripped]
                block_type = "TEXT"
            else:
                current_block.append(stripped)
                    
    if current_block:
        blocks.append((block_type, current_block))
        
    return blocks

def parse_markdown_to_docx(doc, md_file_path):
    with open(md_file_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
    
    lines = md_text.split('\n')
    blocks = get_blocks(lines)
    
    chapter_count = 0
    
    try:
        caption_style = doc.styles['Caption']
    except KeyError:
        caption_style = doc.styles.add_style('Caption', WD_STYLE_TYPE.PARAGRAPH)
    caption_style.font.name = 'Times New Roman'
    caption_style.font.size = Pt(12)
    caption_style.font.bold = True
    caption_style.font.color.rgb = RGBColor(0,0,0)

    for btype, block_lines in blocks:
        if btype == "CODE":
            p = doc.add_paragraph('\n'.join(block_lines))
            p.style = doc.styles['No Spacing']
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
            doc.add_paragraph()
            continue
            
        if btype == "TABLE":
            table_data = []
            for line in block_lines:
                if re.match(r'^\|[\s\-\|]+\|$', line):
                    continue
                cols = [col.strip() for col in line.strip('|').split('|')]
                table_data.append(cols)
                
            if table_data:
                p_cap = doc.add_paragraph(style='Caption')
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_cap = p_cap.add_run("Table ")
                add_seq_field(run_cap, "Table")
                p_cap.add_run(": Data Table")
                p_cap.paragraph_format.space_after = Pt(6)
                
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for r_idx, row in enumerate(table_data):
                    for c_idx, cell_text in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_text
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.name = 'Times New Roman'
                                run.font.size = Pt(11)
                doc.add_paragraph()
            continue
            
        text = " ".join(block_lines).strip()
        
        if chapter_count == 0 and not text.startswith('# '):
            continue
            
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', text)
        if img_match:
            caption_text = img_match.group(1)
            img_path = img_match.group(2)
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img_path, width=Cm(14))
            else:
                p = doc.add_paragraph(f"[Image missing: {img_path}]")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
            p_cap = doc.add_paragraph(style='Caption')
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_cap = p_cap.add_run("Figure ")
            add_seq_field(run_cap, "Figure")
            p_cap.add_run(f": {caption_text}")
            p_cap.paragraph_format.space_before = Pt(6)
            p_cap.paragraph_format.space_after = Pt(12)
            continue
            
        if text.startswith('#'):
            level = len(text) - len(text.lstrip('#'))
            header_text = text.lstrip('#').strip()
            
            is_reference = 'reference' in header_text.lower() or 'appendix' in header_text.lower()
            
            if level == 1 or is_reference:
                if is_reference:
                    level = 1
                    doc.add_page_break()
                    header_text = "References" if 'reference' in header_text.lower() else "Appendix"
                else:
                    chapter_count += 1
                    is_first = (chapter_count == 1)
                    start_chapter(doc, is_first)
                    header_text = f"Chapter-{chapter_count}. {header_text}"
                
            p = doc.add_paragraph(style=f'Heading {level}')
            run = p.add_run(header_text)
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.font.size = Pt(16)
            elif level == 2:
                run.font.size = Pt(14)
            else:
                run.font.size = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            continue
            
        num_match = re.match(r'^(\d+)\.\s(.*)', text)
        if num_match:
            list_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _parse_inline_formatting(p, list_text)
            continue
            
        if text.startswith('- ') or text.startswith('* '):
            list_text = text[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _parse_inline_formatting(p, list_text)
            continue
            
        if text.startswith('> '):
            quote_text = text[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1)
            _parse_inline_formatting(p, quote_text)
            continue
            
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _parse_inline_formatting(p, text)

def generate_report():
    print("[*] Starting Report Generation...")
    doc = initialize_document()
    
    print("[*] Adding Front Matter...")
    add_cover_page(doc)
    start_front_matter(doc)
    add_declaration(doc)
    add_image_page(doc, "OFFER LETTER", "offer_letter.png")
    add_image_page(doc, "CERTIFICATE", "Ifso_certificate.png")
    add_image_page(doc, "EMPLOYER FEEDBACK", "Employer Feedback Form.png")
    add_acknowledgement(doc)
    
    print("[*] Adding TOC, LOF, LOT...")
    add_toc_lof_lot(doc)
    
    print("[*] Parsing Chapters from content.md...")
    parse_markdown_to_docx(doc, "content.md")
    
    output_path = "Final_Report_v4.docx"
    doc.save(output_path)
    print(f"[+] Final Report Complete. Saved to {output_path}")

if __name__ == "__main__":
    generate_report()
