import argparse
import sys

try:
    from docx import Document
except ImportError:
    print("[-] Error: 'python-docx' library is not installed. Please install it using 'pip install python-docx'.")
    sys.exit(1)

def extract_text(docx_path):
    try:
        doc = Document(docx_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    row_data.append(cell.text.replace('\n', ' '))
                full_text.append(" | ".join(row_data))
        
        print('\n'.join(full_text))
    except Exception as e:
        print(f"[-] Failed to read {docx_path}: {e}")
        sys.exit(1)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Extract raw text from a .docx file.")
    parser.add_argument("--input", required=True, help="Path to the .docx file")
    args = parser.parse_args()
    
    extract_text(args.input)
